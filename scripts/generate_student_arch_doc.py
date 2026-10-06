import os
from pathlib import Path
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent.parent
OUT_IMG = ROOT / "docs" / "architecture" / "student_arch_diagram.png"
OUT_DOCX = ROOT / "docs" / "architecture" / "EchoInsight_Master_Architecture.docx"

def draw_bw_diagram():
    fig, ax = plt.subplots(figsize=(11, 8.5)) # US Letter landscape
    ax.set_facecolor("white")
    fig.patch.set_facecolor("white")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    def _box(x, y, w, h, label, sub="", style="solid", lw=1.5):
        ls = "dashed" if style == "dashed" else "solid"
        box = mpatches.FancyBboxPatch((x - w/2, y - h/2), w, h,
                                      boxstyle="round,pad=0.02",
                                      linewidth=lw, edgecolor="black", 
                                      facecolor="white", linestyle=ls)
        ax.add_patch(box)
        if sub:
            ax.text(x, y + 0.02, label, ha="center", va="center", fontsize=9, fontweight="bold", color="black")
            ax.text(x, y - 0.04, sub, ha="center", va="center", fontsize=7, color="black", style="italic")
        else:
            ax.text(x, y, label, ha="center", va="center", fontsize=9, fontweight="bold", color="black")

    def _arrow(x1, y1, x2, y2, label=""):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="->", color="black", lw=1.2, connectionstyle="arc3,rad=0.0"))
        if label:
            mx, my = (x1 + x2) / 2, (y1 + y2) / 2
            ax.text(mx, my+0.01, label, fontsize=7, color="black", ha="center")

    # Title
    ax.text(0.5, 0.96, "EchoInsight - Master Architecture Diagram", ha="center", fontsize=14, fontweight="bold")
    
    # Users
    _box(0.1, 0.8, 0.15, 0.08, "Users & Roles", "Admin, Supervisor, Agent")
    
    # Frontend
    _box(0.1, 0.6, 0.16, 0.12, "Frontend (Vercel)", "React 18, Vite, Tailwind\nSPA, TanStack Query")
    
    # Backend Boundary
    bg = mpatches.Rectangle((0.25, 0.3), 0.5, 0.6, fill=False, edgecolor="black", linestyle="dotted", lw=1)
    ax.add_patch(bg)
    ax.text(0.26, 0.88, "Backend API & Worker (Render.com Docker)", fontsize=9, fontweight="bold")
    
    # Ingest
    _box(0.35, 0.75, 0.16, 0.1, "Ingest & Redact", "Auth (JWT, RBAC)\n9 PII Regex Patterns")
    
    # Redaction Boundary
    ax.plot([0.45, 0.45], [0.35, 0.85], color="black", lw=2)
    ax.text(0.44, 0.6, "REDACTION BOUNDARY", rotation=90, va="center", ha="right", fontsize=8, fontweight="bold")
    
    # Pipeline & Worker
    _box(0.55, 0.75, 0.16, 0.1, "Core Pipeline", "State Reducer\nCommitment Ledger")
    _box(0.55, 0.60, 0.16, 0.08, "Background Worker", "Polls jobs every 5s")
    
    # Gates
    _box(0.55, 0.45, 0.16, 0.1, "Validation Gates", "Evidence Gate (exact match)\nHard Gate (7 checks)\nPhrase Matcher\nQA Scorer (60% cap)")
    
    # LLM
    _box(0.88, 0.60, 0.18, 0.15, "Intelligence Layer", "Primary: OpenRouter\nFallback: Groq\nModels: qwen3.8-27b,\ngpt-oss-20b\n(Redacted text only)")
    
    # DB
    _box(0.88, 0.40, 0.16, 0.12, "Data Stores", "PostgreSQL (20 tables)\ndemo_seed.db (SQLite)\nYAML configs")
    
    # Optional
    _box(0.88, 0.8, 0.16, 0.1, "Optional Layers", "Action Layer (Off)\nChat Assistant (Active)", style="dashed")
    
    # Foundation
    _box(0.5, 0.15, 0.8, 0.1, "Foundation & Operations", "Trust: RBAC, Audit logging, Prompt injection defense\nObservability: /health, /metrics, Token budget\nDeploy: Render Free Tier (30-60s cold start)")

    # Connections
    _arrow(0.1, 0.76, 0.1, 0.66, "")
    _arrow(0.18, 0.6, 0.27, 0.75, "HTTPS / JWT")
    _arrow(0.43, 0.75, 0.47, 0.75, "Redacted")
    _arrow(0.55, 0.7, 0.55, 0.64, "Queues job")
    _arrow(0.63, 0.6, 0.79, 0.6, "SYSTEM_TURN\nSYSTEM_FINAL")
    _arrow(0.79, 0.58, 0.63, 0.48, "JSON outputs")
    _arrow(0.63, 0.45, 0.8, 0.45, "Verified Writes")
    
    fig.savefig(str(OUT_IMG), format="png", dpi=200, bbox_inches="tight")
    plt.close(fig)

def create_docx():
    doc = Document()
    
    # Set to landscape
    section = doc.sections[0]
    new_width, new_height = section.page_height, section.page_width
    section.page_width = new_width
    section.page_height = new_height

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("EchoInsight Architecture")
    run.bold = True
    run.font.size = Pt(16)
    
    # Add Diagram
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(OUT_IMG), width=Inches(9.0))
    
    # Add text summary to make it structured and student-like
    doc.add_heading('Architecture Summary', level=1)
    
    p = doc.add_paragraph()
    p.add_run("1. Frontend (Vercel): ").bold = True
    p.add_run("React 18 SPA using Vite and Tailwind. Connects to backend via HTTPS with short-lived JWTs. Includes views for Dashboard, Conversation Detail, and Admin Panel.")
    
    p = doc.add_paragraph()
    p.add_run("2. Backend API (Render.com Docker): ").bold = True
    p.add_run("FastAPI Modular Monolith. Contains the Core Pipeline (Ingest, State Reducer, Commitment Ledger). A strict REDACTION BOUNDARY uses 9 regex patterns to remove PII before data is stored or sent to the LLM.")
    
    p = doc.add_paragraph()
    p.add_run("3. Validation Gates: ").bold = True
    p.add_run("Prevents hallucinations. Includes an Evidence Gate (exact substring match), a Hard Gate (structural checks), and a QA Scorer (with a 60% critical violation cap).")
    
    p = doc.add_paragraph()
    p.add_run("4. Intelligence Layer: ").bold = True
    p.add_run("Primary provider is OpenRouter, with Groq as fallback. Uses qwen3.8-27b for extraction and gpt-oss-20b for selective verification. Only ever receives redacted text.")
    
    p = doc.add_paragraph()
    p.add_run("5. Data Stores: ").bold = True
    p.add_run("PostgreSQL for production. Uses demo_seed.db (SQLite) restored at boot if the DB is empty on Render's ephemeral free tier.")
    
    doc.save(str(OUT_DOCX))
    print(f"Generated {OUT_DOCX}")

if __name__ == "__main__":
    draw_bw_diagram()
    create_docx()
