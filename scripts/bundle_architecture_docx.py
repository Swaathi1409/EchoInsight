"""
Bundles the generated diagrams into a Word document.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).parent.parent
DIAGRAMS_DIR = ROOT / "docs" / "architecture" / "diagrams"
OUTPUT_DOCX = ROOT / "docs" / "architecture" / "EchoInsight_Architecture_Diagram_Pack.docx"

def bundle_diagrams():
    doc = Document()
    doc.add_heading("EchoInsight Architecture Diagram Pack", 0)
    
    diagrams = [
        ("D0 Master Poster", "D0_master_poster.png", "System overview covering deployment, observability, and data flow."),
        ("D1 System Context", "D1_system_context.png", "System boundaries, users, and external dependencies."),
        ("D2 Runtime Containers", "D2_runtime_containers.png", "Runtime processes, files, and stores."),
        ("D3 Backend Components", "D3_backend_components.png", "Python packages and import directions."),
        ("D4 Analysis Pipeline", "D4_analysis_pipeline.png", "Incremental and final analysis data flows."),
        ("D5 State Machine", "D5_state_machine.png", "Conversation lifecycle states and transitions."),
        ("D6 Commitment Ledger", "D6_commitment_ledger.png", "Commitment status transitions."),
        ("D7 QA Scoring Flow", "D7_qa_scoring.png", "QA pipeline, selective verification, and scoring formula."),
        ("D8 Sequence Diagrams", "D8_sequence_diagrams.png", "Operation order for turn append and final analysis."),
        ("D9 Data Model", "D9_data_model.png", "ER diagram for the PostgreSQL schema."),
        ("D10 Security Boundaries", "D10_security.png", "PII redaction and access scopes."),
        ("D11 Deployment Topology", "D11_deployment.png", "Vercel, Render, and cold start behavior."),
        ("D12 Observability & CI", "D12_observability.png", "Metrics, health endpoints, and eval pipeline."),
        ("D13 Frontend Architecture", "D13_frontend.png", "React SPA structure and component routing."),
        ("D14 Optional Layers", "D14_optional_layers.png", "Action layer and assistant integrations."),
        ("D15 Scale-Out Target", "D15_scale_out.png", "Future production scale-out design.")
    ]

    for title, filename, caption in diagrams:
        img_path = DIAGRAMS_DIR / filename
        if img_path.exists():
            doc.add_heading(title, level=1)
            doc.add_paragraph(caption)
            doc.add_picture(str(img_path), width=Inches(6.0))
            doc.add_page_break()
        else:
            print(f"Warning: Missing {filename}")

    doc.save(str(OUTPUT_DOCX))
    print(f"Saved {OUTPUT_DOCX}")

if __name__ == "__main__":
    bundle_diagrams()
